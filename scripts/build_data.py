#!/usr/bin/env python3
"""Build Cleanest Hour's static data files from public U.S. government sources.

Inputs (downloaded by hand, not committed; see README "Rebuilding the data"):
  * EIA-930 Hourly Electric Grid Monitor six-month "BALANCE" CSVs
    https://www.eia.gov/electricity/gridmonitor/sixMonthFiles/EIA930_BALANCE_<YYYY>_<Mon>_<Mon>.csv
  * EPA eGRID2023 data file, revision 2
    https://www.epa.gov/system/files/documents/2025-06/egrid2023_data_rev2.xlsx

Outputs (committed, served statically):
  data/index.json         manifest: regions, factors, sources, validation numbers
  data/<BA>.json          one file per balancing authority: 365 days x 24 hourly
                          carbon-intensity values (g CO2e/kWh) + hourly fuel mix

Method (also explained in the app's "How this is calculated" panel):
  intensity(hour) = sum_f gen_f(hour) * rate_f / sum_f gen_f(hour)
  where gen_f is EIA-930 *adjusted* net generation by energy source (MW, i.e. MWh
  for the hour), negative values (storage charging, plant self-use) are set to 0,
  and rate_f is an EPA eGRID2023 CO2-equivalent output emission rate:
    coal, gas  -> the balancing authority's own eGRID coal / gas rate
    oil        -> BA oil rate if 1000-3500 lb/MWh, else the U.S. oil rate
    other/unknown fuel -> the BA's eGRID annual average (total) output rate
    nuclear, hydro, wind, solar, geothermal, storage discharge -> 0 (no combustion)

Usage:
  python3 -I scripts/build_data.py --eia RAW/EIA930_BALANCE_*.csv --egrid RAW/egrid2023_data_rev2.xlsx \
      --start 2025-10-01 --end 2026-09-30 --out data
"""
import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import re
import sys
from zoneinfo import ZoneInfo

LB_PER_MWH_TO_G_PER_KWH = 0.45359237  # 1 lb = 453.59237 g; 1 MWh = 1000 kWh

# Balancing authorities shown in the app, the civil time zone used to display hours
# (most of the BA's customers), and a plain-language description of the area served.
REGIONS = [
    ("PJM", "America/New_York", "PJM Interconnection", "NJ, PA, MD, DE, DC, VA, WV, OH, and parts of IL (Chicago area), IN, MI, KY, NC"),
    ("MISO", "America/Chicago", "Midcontinent ISO", "MN, IA, WI, most of MI, IL (outside Chicago), IN, MO, AR, LA, MS, ND, parts of TX, MT, SD, KY"),
    ("ERCO", "America/Chicago", "ERCOT (Texas)", "Most of Texas"),
    ("CISO", "America/Los_Angeles", "California ISO", "Most of California (PG&E, SCE, SDG&E areas)"),
    ("SWPP", "America/Chicago", "Southwest Power Pool", "KS, OK, NE, parts of TX panhandle, NM, AR, LA, MO, SD, ND, MT, WY"),
    ("NYIS", "America/New_York", "New York ISO", "New York State"),
    ("ISNE", "America/New_York", "ISO New England", "MA, CT, RI, NH, VT, most of ME"),
    ("SOCO", "America/New_York", "Southern Company", "Georgia, Alabama, Mississippi (Southern Co. utilities)"),
    ("FPL", "America/New_York", "Florida Power & Light", "Southern and eastern Florida"),
    ("FPC", "America/New_York", "Duke Energy Florida", "Central and northern Florida"),
    ("TVA", "America/Chicago", "Tennessee Valley Authority", "Tennessee and parts of AL, MS, KY, GA, NC, VA"),
    ("DUK", "America/New_York", "Duke Energy Carolinas", "Western North Carolina and western South Carolina"),
    ("CPLE", "America/New_York", "Duke Energy Progress East", "Eastern North Carolina and northeastern South Carolina"),
    ("BPAT", "America/Los_Angeles", "Bonneville Power Administration", "Much of Washington, Oregon, Idaho (Pacific Northwest)"),
    ("PACW", "America/Los_Angeles", "PacifiCorp West", "Pacific Power areas of OR, WA, northern CA"),
    ("PACE", "America/Denver", "PacifiCorp East", "Utah, Wyoming, southeastern Idaho (Rocky Mountain Power)"),
    ("PSCO", "America/Denver", "Public Service Co. of Colorado", "Most of Colorado (Xcel Energy)"),
    ("AZPS", "America/Phoenix", "Arizona Public Service", "Much of Arizona outside Phoenix's SRP area"),
    ("SRP", "America/Phoenix", "Salt River Project", "Phoenix metro (SRP customers)"),
    ("NEVP", "America/Los_Angeles", "Nevada Power", "Southern Nevada (Las Vegas)"),
    ("LDWP", "America/Los_Angeles", "LA Dept. of Water & Power", "City of Los Angeles"),
    ("BANC", "America/Los_Angeles", "Balancing Authority of Northern California", "Sacramento area (SMUD) and parts of northern California"),
    ("PGE", "America/Los_Angeles", "Portland General Electric", "Portland, Oregon area"),
    ("PSEI", "America/Los_Angeles", "Puget Sound Energy", "Puget Sound area of Washington"),
    ("IPCO", "America/Denver", "Idaho Power", "Southern Idaho and eastern Oregon"),
    ("NWMT", "America/Denver", "NorthWestern Energy (Montana)", "Most of Montana"),
    ("PNM", "America/Denver", "Public Service Co. of New Mexico", "Much of New Mexico"),
    ("TEPC", "America/Phoenix", "Tucson Electric Power", "Tucson and southeastern Arizona"),
    ("SCEG", "America/New_York", "Dominion Energy South Carolina", "Central and southern South Carolina"),
    ("SC", "America/New_York", "Santee Cooper", "Parts of South Carolina"),
    ("LGEE", "America/New_York", "LG&E and KU", "Most of Kentucky"),
]

TZ_LABEL = {
    "America/New_York": "Eastern time",
    "America/Chicago": "Central time",
    "America/Denver": "Mountain time",
    "America/Phoenix": "Arizona time (no DST)",
    "America/Los_Angeles": "Pacific time",
}

# Fuel groups. Keys are matched against normalised EIA column names.
FUEL_COLUMNS = {
    "coal": ["coal"],
    "gas": ["natural gas"],
    "oil": ["all petroleum products"],
    "nuclear": ["nuclear"],
    "hydro": ["hydropower excluding pumped storage"],
    "wind": ["wind without integrated battery storage", "wind with integrated battery storage"],
    "solar": ["solar without integrated battery storage", "solar with integrated battery storage",
              "solar witho integrated battery storage"],  # EIA header typo in 2025-26 files
    "geothermal": ["geothermal"],
    "storage": ["battery storage", "pumped storage", "other energy storage", "unknown energy storage"],
    "other": ["other fuel sources", "unknown fuel sources"],
}
FUELS = list(FUEL_COLUMNS)  # order used in the mix arrays
ZERO_COMBUSTION = {"nuclear", "hydro", "wind", "solar", "geothermal", "storage"}
OIL_BAND = (1000.0, 3500.0)  # lb/MWh; outside this the BA oil rate is treated as unreliable
SEASONS = {12: "winter", 1: "winter", 2: "winter", 3: "spring", 4: "spring", 5: "spring",
           6: "summer", 7: "summer", 8: "summer", 9: "fall", 10: "fall", 11: "fall"}


def norm(s):
    return re.sub(r"\s+", " ", s.strip().lower())


def map_columns(header):
    """Return {fuel: [column indexes]} for the *adjusted* by-source columns, plus total NG."""
    out = {f: [] for f in FUELS}
    total_ng = None
    for i, raw in enumerate(header):
        h = norm(raw)
        if h == "net generation (mw) (adjusted)":
            total_ng = i
            continue
        m = re.match(r"net generation \(mw\) from (.+) \(adjusted\)$", h)
        if not m:
            continue
        src = m.group(1).strip()
        for fuel, names in FUEL_COLUMNS.items():
            if src in names:
                out[fuel].append(i)
                break
        else:
            sys.exit(f"Unmapped EIA fuel column: {raw!r}")
    missing = [f for f, ix in out.items() if not ix]
    if missing or total_ng is None:
        sys.exit(f"Missing columns {missing} / total NG {total_ng}")
    return out, total_ng


def num(s):
    if s == "" or s is None:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def load_egrid(path):
    import openpyxl  # build-time only dependency

    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb["BA23"].iter_rows(values_only=True))
    codes = rows[1]
    ix = {k: i for i, k in enumerate(codes)}
    us_rows = list(wb["US23"].iter_rows(values_only=True))
    us = dict(zip(us_rows[1], us_rows[2]))
    out = {}
    for r in rows[2:]:
        code = r[ix["BACODE"]]
        if not code:
            continue
        g = lambda k: r[ix[k]]
        out[code] = {
            "name": g("BANAME"),
            "coal": g("BACC2ERT"), "oil": g("BAOC2ERT"), "gas": g("BAGC2ERT"),
            "total": g("BAC2ERTA"),
        }
    return out, {"coal": us["USCC2ERT"], "oil": us["USOC2ERT"], "gas": us["USGC2ERT"], "total": us["USC2ERTA"]}


def factors_for(ba, egrid, us):
    e = egrid[ba]
    notes = []
    lb = {}
    for fuel in ("coal", "gas"):
        v = e[fuel]
        if v is None or v <= 0:
            v = us[fuel]
            notes.append(f"{fuel}: BA rate missing, used U.S. rate")
        lb[fuel] = float(v)
    oil = e["oil"]
    if oil is None or not (OIL_BAND[0] <= oil <= OIL_BAND[1]):
        notes.append(f"oil: BA rate {oil!r} lb/MWh outside {OIL_BAND}, used U.S. oil rate")
        oil = us["oil"]
    lb["oil"] = float(oil)
    lb["other"] = float(e["total"])
    for f in ZERO_COMBUSTION:
        lb[f] = 0.0
    g = {f: round(v * LB_PER_MWH_TO_G_PER_KWH, 1) for f, v in lb.items()}
    return g, {f: round(v, 1) for f, v in lb.items()}, notes


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--eia", nargs="+", required=True)
    ap.add_argument("--egrid", required=True)
    ap.add_argument("--start", required=True)  # inclusive local date
    ap.add_argument("--end", required=True)    # inclusive local date
    ap.add_argument("--out", required=True)
    ap.add_argument("--lastmod", default="{}", help="JSON {filename: Last-Modified header}")
    args = ap.parse_args()

    start = dt.date.fromisoformat(args.start)
    end = dt.date.fromisoformat(args.end)
    ndays = (end - start).days + 1
    lastmod = json.loads(args.lastmod)

    egrid, us = load_egrid(args.egrid)
    region_codes = {r[0] for r in REGIONS}
    tzs = {r[0]: ZoneInfo(r[1]) for r in REGIONS}

    # per BA: hourly slots [day][hour] -> dict fuel->MWh (first occurrence wins on DST fall-back)
    slots = {ba: [[None] * 24 for _ in range(ndays)] for ba in region_codes}
    qa = {ba: {"rows": 0, "sum_fuels": 0.0, "sum_total": 0.0, "dup_hours": 0} for ba in region_codes}

    for path in args.eia:
        with open(path, newline="") as f:
            rd = csv.reader(f)
            header = next(rd)
            cols, total_ix = map_columns(header)
            for row in rd:
                ba = row[0]
                if ba not in region_codes:
                    continue
                # UTC time at END of hour -> start of hour -> display time zone
                t_end = dt.datetime.strptime(row[4], "%m/%d/%Y %I:%M:%S %p").replace(tzinfo=dt.timezone.utc)
                t0 = (t_end - dt.timedelta(hours=1)).astimezone(tzs[ba])
                d = (t0.date() - start).days
                if d < 0 or d >= ndays:
                    continue
                gen = {}
                raw_sum = 0.0
                any_val = False
                for fuel, ixs in cols.items():
                    s = 0.0
                    for i in ixs:
                        v = num(row[i])
                        if v is not None:
                            any_val = True
                            raw_sum += v
                            s += max(v, 0.0)  # charging / self-use -> 0
                    gen[fuel] = s
                if not any_val:
                    continue
                if slots[ba][d][t0.hour] is not None:
                    qa[ba]["dup_hours"] += 1  # DST fall-back repeated hour; keep first
                    continue
                slots[ba][d][t0.hour] = gen
                q = qa[ba]
                q["rows"] += 1
                q["sum_fuels"] += raw_sum  # unclamped, comparable to EIA total NG
                tot = num(row[total_ix])
                if tot is not None:
                    q["sum_total"] += tot

    os.makedirs(args.out, exist_ok=True)
    index_regions = []
    for code, tzname, label, area in REGIONS:
        q = qa[code]
        coverage = q["rows"] / (ndays * 24)
        ratio = q["sum_fuels"] / q["sum_total"] if q["sum_total"] else 0
        if coverage < 0.95 or not (0.9 <= ratio <= 1.1):
            print(f"DROP {code}: coverage {coverage:.3f}, fuel/total ratio {ratio:.3f}", file=sys.stderr)
            continue
        g, lb, notes = factors_for(code, egrid, us)
        ci = []
        mix_acc = {s: [[0.0] * len(FUELS) for _ in range(24)] for s in ("all", "winter", "spring", "summer", "fall")}
        em_total = 0.0
        gen_total = 0.0
        fuel_totals = [0.0] * len(FUELS)
        nulls = 0
        for d in range(ndays):
            day = start + dt.timedelta(days=d)
            season = SEASONS[day.month]
            out_day = []
            for h in range(24):
                gen = slots[code][d][h]
                tot = sum(gen.values()) if gen else 0.0
                if not gen or tot <= 0:
                    out_day.append(None)
                    nulls += 1
                    continue
                em = sum(gen[f] * g[f] for f in FUELS)  # MWh * g/kWh = kg ... ratio is what matters
                out_day.append(round(em / tot))
                em_total += em
                gen_total += tot
                for k, f in enumerate(FUELS):
                    fuel_totals[k] += gen[f]
                    share = gen[f] / tot
                    mix_acc["all"][h][k] += share
                    mix_acc[season][h][k] += share
            ci.append(out_day)
        # average hourly shares -> permille ints
        mix = {}
        for s, rows in mix_acc.items():
            mix[s] = []
            for h in range(24):
                ssum = sum(rows[h]) or 1
                mix[s].append([round(1000 * v / ssum) for v in rows[h]])
        annual_avg = em_total / gen_total
        egrid_avg = egrid[code]["total"] * LB_PER_MWH_TO_G_PER_KWH
        file_obj = {
            "ba": code, "name": label, "area": area, "tz": tzname, "tzLabel": TZ_LABEL[tzname],
            "start": args.start, "days": ndays, "fuels": FUELS,
            "factors_g_per_kwh": g,
            "ci": ci, "mix": mix,
        }
        with open(os.path.join(args.out, f"{code}.json"), "w") as f:
            json.dump(file_obj, f, separators=(",", ":"))
        index_regions.append({
            "ba": code, "name": label, "area": area, "tz": tzname, "tzLabel": TZ_LABEL[tzname],
            "annualAvg_g_per_kwh": round(annual_avg, 1),
            "profileAll": [round(sum(v[h] for v in ci if v[h] is not None) / max(1, sum(1 for v in ci if v[h] is not None))) for h in range(24)],
            "egrid2023Avg_g_per_kwh": round(egrid_avg, 1),
            "coverage": round(coverage, 4), "nullHours": nulls, "dupHoursDropped": q["dup_hours"],
            "fuelToTotalRatio": round(ratio, 4),
            "genShare": {f: round(fuel_totals[k] / gen_total, 4) for k, f in enumerate(FUELS)},
            "factors_lb_per_mwh": lb, "factorNotes": notes,
        })
        print(f"{code:5s} cov {coverage:.3f} ratio {ratio:.3f} nulls {nulls:4d} avg {annual_avg:6.1f} g/kWh  eGRID23 {egrid_avg:6.1f}  {'; '.join(notes)}")

    index = {
        "generated": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "dateRange": {"start": args.start, "end": args.end, "days": ndays},
        "units": "g CO2e per kWh generated; hour h = clock hour starting at h:00 in the region's display time zone",
        "fuels": FUELS,
        "sources": {
            "eia930": [{"file": os.path.basename(p), "bytes": os.path.getsize(p), "sha256": sha256(p),
                        "lastModified": lastmod.get(os.path.basename(p))} for p in args.eia],
            "egrid": {"file": os.path.basename(args.egrid), "sha256": sha256(args.egrid),
                      "url": "https://www.epa.gov/system/files/documents/2025-06/egrid2023_data_rev2.xlsx"},
        },
        "usFactors_lb_per_mwh": {k: round(float(v), 1) for k, v in us.items()},
        "regions": index_regions,
    }
    with open(os.path.join(args.out, "index.json"), "w") as f:
        json.dump(index, f, indent=1)


if __name__ == "__main__":
    main()
